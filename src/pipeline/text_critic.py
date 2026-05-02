"""Stage 5-A — per-slide text critique."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

from src.llm.ollama_client import chat, DEFAULT_MODEL
from src.pipeline.planner import parse_json_block

PROMPTS_DIR = Path(__file__).resolve().parents[2] / "prompts"


def extract_slide_text(pptx_path: str | Path) -> list[str]:
    """Read all non-empty text frame contents from the first slide of pptx_path."""
    prs = Presentation(str(pptx_path))
    if not prs.slides:
        return []
    slide = prs.slides[0]
    texts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                texts.append(t)
    return texts


def critique_slide_text(
    slide_plan: dict,
    pptx_path: str | Path,
    *,
    model: str = DEFAULT_MODEL,
) -> dict:
    """Ask LLM if the rendered text matches the plan. Returns {issues, verdict}."""
    texts = extract_slide_text(pptx_path)
    extracted = "\n".join(f"- {t}" for t in texts) if texts else "(no text on slide)"
    template = (PROMPTS_DIR / "text_critique.txt").read_text(encoding="utf-8")
    prompt = template.format(
        slide_plan_json=json.dumps(slide_plan, ensure_ascii=False, indent=2),
        extracted_text=extracted,
    )
    response = chat(prompt, model=model)
    parsed = parse_json_block(response)
    if not isinstance(parsed, dict):
        raise ValueError("expected JSON object for text critique")
    parsed.setdefault("issues", [])
    parsed.setdefault("verdict", "PASS")
    return parsed
