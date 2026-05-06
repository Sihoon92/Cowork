"""Stage 0–4 orchestrator. Stages 5–7 are added in later tasks.

See spec §4.
"""
from __future__ import annotations

import json
import time
from pathlib import Path
from typing import Optional

from src.pipeline_v2.auto_fix import execute_with_auto_fix
from src.pipeline_v2.code_assembler import PREAMBLE_PATH
from src.pipeline_v2.code_generator import generate_slide_body
from src.pipeline_v2.outline import generate_outline
from src.pipeline_v2.visual_strategy import add_visual_strategy
from src.pipeline_v2.gallery import Gallery
from src.pptx.merger import merge_slides
from src.util import log


GUIDELINES_DIR = Path(__file__).resolve().parents[2] / "guidelines" / "v2"


def build_presentation_v2(
    content: dict,
    *,
    output_path: Path,
    workdir: Path,
    gallery: Optional[Gallery] = None,
    enable_revision: bool = True,
    enable_gallery: bool = True,
    max_iter: int = 3,
) -> Path:
    t0 = time.time()
    output_path = Path(output_path)
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    slides_dir = workdir / "slides"; slides_dir.mkdir(exist_ok=True)
    code_dir = workdir / "slide_codes"; code_dir.mkdir(exist_ok=True)

    # Stage 1a + 1b
    outline = generate_outline(content)
    (workdir / "outline.json").write_text(
        json.dumps(outline, ensure_ascii=False, indent=2), encoding="utf-8")

    plan = add_visual_strategy(
        outline,
        gallery=gallery if (enable_gallery and gallery is not None) else None,
    )
    (workdir / "plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")

    # Stage 2 + 3 (per-slide LLM codegen + auto-fix execution)
    design_guide = (GUIDELINES_DIR / "design_guide.md").read_text(encoding="utf-8")
    preamble_text = PREAMBLE_PATH.read_text(encoding="utf-8")

    slide_paths: list[Path] = []
    for s in plan["slides"]:
        idx = s["index"]
        log.stage(f"v2 Stage 2/3: slide {idx} ({s['kind']})")
        body = generate_slide_body(
            s, design_guide=design_guide, preamble_excerpt=preamble_text)
        out_pptx = slides_dir / f"slide_{idx:02d}.pptx"
        try:
            execute_with_auto_fix(
                body, out_pptx, work_dir=code_dir,
                design_guide_excerpt=design_guide[:2000],
            )
        except Exception as exc:  # noqa: BLE001
            log.warn(f"slide {idx} failed after retries: {exc}; placeholder used")
            _write_blank_slide(out_pptx, reason=str(exc))
        slide_paths.append(out_pptx)

    # Stage 4 — merge
    log.stage("v2 Stage 4: merge")
    merge_slides(slide_paths, output_path)
    log.ok(f"-> {output_path}  ({output_path.stat().st_size // 1024} KB)")

    # Stages 5–7 added in Tasks 9 & 10
    log.ok(f"v2 done in {time.time() - t0:.1f}s")
    return output_path


def _write_blank_slide(out_pptx: Path, *, reason: str) -> None:
    """Last-resort fallback so merge does not break."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(2))
    tb.text_frame.text = f"slide failed: {reason[:200]}"
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))
