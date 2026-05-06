"""Integration tests for the LLM-codegen rendering path.

The actual LLM call is mocked. We exercise the glue between
`code_generator`, `code_runner`, and the per-slide error fallback so the
codegen path is verifiable without a live model.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from src.pipeline import codegen_path
from src.pptx.code_runner import CodeExecutionError


# A minimal valid script the runtime can execute end-to-end. The codegen
# path normally returns code wrapped in a ```python``` fence; the
# `extract_code_block` step strips that, so we bypass the fence here.
_VALID_SCRIPT = '''
def add_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    deck_meta = data["deck_meta"]
    body = apply_master(slide, deck_meta,
                        head_message=data["head_message"],
                        page_no=data.get("slide_no"),
                        section_id=data.get("section_id"))
    add_text(slide, rect=body, text=data["head_message"],
             font_size=20, color="#000000")
'''

_BROKEN_SCRIPT = '''
def add_slide(prs, data):
    raise RuntimeError("intentionally broken for retry test")
'''


def _make_flat_slide(*, intent_label: str = "single_metric_emphasis",
                    layout_hint: str = "test layout") -> dict:
    return {
        "slide_no": 2,
        "section_id": "scale",
        "head_message": "테스트 헤드",
        "intent_label": intent_label,
        "visual_strategy": {
            "approach": "test_approach",
            "layout_hint": layout_hint,
            "key_elements": ["a", "b"],
        },
        "_full": {
            "content": {
                "slide_no": 2,
                "intent_label": intent_label,
                "head_message": "테스트 헤드",
                "key_takeaway": "테스트 인사이트",
                "knowledge": {"metrics": [], "narratives": []},
            }
        },
    }


def _deck_meta() -> dict:
    return {
        "title": "test deck",
        "theme": {
            "background": "#FFFFFF", "head_text": "#1C2833",
            "head_rule": "#065A82", "body_text": "#1C2833",
            "bullet": "#065A82", "accent": "#21295C",
            "primary": "#065A82", "neutral": "#F2F2F2",
            "text_dark": "#1C2833", "text_light": "#FFFFFF",
        },
        "fonts": {"header": "Pretendard Bold", "body": "Pretendard"},
        "slide_size": {"width_in": 13.333, "height_in": 7.5},
        "section_titles": {"scale": "Scale"},
    }


def test_render_via_codegen_happy_path(tmp_path: Path, monkeypatch):
    """Mocked LLM returns a working script — output .pptx is produced."""
    calls: list[dict] = []

    def fake_generate(*, layout_hint, pattern_guideline, slide_data, model):
        calls.append({
            "layout_hint": layout_hint,
            "slide_data_keys": sorted(slide_data.keys()),
        })
        return _VALID_SCRIPT

    monkeypatch.setattr(codegen_path, "generate_slide_code", fake_generate)

    out = tmp_path / "slide_02.pptx"
    codegen_path.render_via_codegen(_deck_meta(), _make_flat_slide(), out)

    assert out.exists()
    prs = Presentation(out)
    assert len(prs.slides) == 1

    # Verify that what we sent to the LLM stub matches the contract
    assert len(calls) == 1
    assert calls[0]["layout_hint"] == "test layout"
    assert "deck_meta" in calls[0]["slide_data_keys"]
    assert "content" in calls[0]["slide_data_keys"]
    assert "head_message" in calls[0]["slide_data_keys"]


def test_render_via_codegen_skips_framing_slides(tmp_path: Path):
    """Cover/closing/divider slides must NOT go through codegen."""
    flat = _make_flat_slide(intent_label="deck_opening")
    out = tmp_path / "cover.pptx"
    with pytest.raises(CodeExecutionError, match="codegen path skipped"):
        codegen_path.render_via_codegen(_deck_meta(), flat, out)


def test_render_via_codegen_falls_back_to_intent_hint(tmp_path: Path, monkeypatch):
    """Missing visual_strategy → use intent-derived fallback hint."""
    captured: dict = {}

    def fake_generate(*, layout_hint, pattern_guideline, slide_data, model):
        captured["layout_hint"] = layout_hint
        return _VALID_SCRIPT

    monkeypatch.setattr(codegen_path, "generate_slide_code", fake_generate)

    flat = _make_flat_slide()
    flat["visual_strategy"] = None  # simulate planner skipped Phase 1b

    out = tmp_path / "slide.pptx"
    codegen_path.render_via_codegen(_deck_meta(), flat, out)

    # Should have used the intent-derived fallback for single_metric_emphasis
    assert captured["layout_hint"] != "test layout"
    assert "metric" in captured["layout_hint"] or "hero" in captured["layout_hint"]


def test_render_via_codegen_retries_on_broken_script(tmp_path: Path, monkeypatch):
    """First attempt explodes at runtime → fix_callback is invoked → second attempt works."""
    call_count = {"n": 0}

    def fake_generate(*, layout_hint, pattern_guideline, slide_data, model):
        call_count["n"] += 1
        return _BROKEN_SCRIPT if call_count["n"] == 1 else _VALID_SCRIPT

    monkeypatch.setattr(codegen_path, "generate_slide_code", fake_generate)

    out = tmp_path / "slide.pptx"
    codegen_path.render_via_codegen(_deck_meta(), _make_flat_slide(), out)

    assert call_count["n"] == 2  # one initial + one retry
    assert out.exists()


def test_intent_fallback_hint_known_intents_have_concrete_directives():
    """Sanity-check the table doesn't degenerate to a single generic line."""
    hints = {
        intent: codegen_path._intent_fallback_hint(intent)
        for intent in (
            "single_metric_emphasis", "multi_metric_dashboard",
            "numeric_comparison", "two_dim_compare", "before_after",
            "sequence_or_timeline", "quotation", "parallel_compare",
            "general_facts",
        )
    }
    # No duplicates: each intent should get a distinct hint
    assert len(set(hints.values())) == len(hints)
    # Every hint mentions either pt-size or a layout primitive (콘크리트성)
    for intent, hint in hints.items():
        assert "pt" in hint or "그리드" in hint or "분할" in hint or "bullet" in hint or "chevron" in hint, (
            f"hint for {intent} too vague: {hint}"
        )


def test_build_slide_data_includes_required_keys():
    flat = _make_flat_slide()
    deck_meta = _deck_meta()
    data = codegen_path._build_slide_data(deck_meta, flat)

    # The code_generation.md prompt teaches the model to read these:
    assert data["deck_meta"] is deck_meta
    assert data["head_message"] == "테스트 헤드"
    assert data["slide_no"] == 2
    assert data["content"]["intent_label"] == "single_metric_emphasis"
    assert data["visual_strategy"]["approach"] == "test_approach"


def test_render_via_codegen_persists_code_artifact(tmp_path: Path, monkeypatch):
    """When code_dir is provided, every attempt is dumped to disk for audit."""
    monkeypatch.setattr(
        codegen_path, "generate_slide_code",
        lambda *, layout_hint, pattern_guideline, slide_data, model: _VALID_SCRIPT,
    )

    code_dir = tmp_path / "codegen"
    out = tmp_path / "slide_02.pptx"
    codegen_path.render_via_codegen(
        _deck_meta(), _make_flat_slide(), out, code_dir=code_dir,
    )

    expected = code_dir / "slide_02_v1.py"
    assert expected.exists(), f"missing artifact: {list(code_dir.glob('*'))}"
    text = expected.read_text(encoding="utf-8")
    assert "def add_slide" in text


def test_render_via_codegen_persists_retry_artifacts(tmp_path: Path, monkeypatch):
    """Retry attempts each get their own _vN_retry artifact file."""
    call_count = {"n": 0}

    def fake_generate(*, layout_hint, pattern_guideline, slide_data, model):
        call_count["n"] += 1
        return _BROKEN_SCRIPT if call_count["n"] == 1 else _VALID_SCRIPT

    monkeypatch.setattr(codegen_path, "generate_slide_code", fake_generate)

    code_dir = tmp_path / "codegen"
    out = tmp_path / "slide_02.pptx"
    codegen_path.render_via_codegen(
        _deck_meta(), _make_flat_slide(), out, code_dir=code_dir,
    )

    artifacts = sorted(p.name for p in code_dir.glob("slide_02_*.py"))
    assert artifacts == ["slide_02_v1.py", "slide_02_v2_retry.py"], artifacts


def test_render_via_codegen_no_code_dir_skips_persistence(tmp_path: Path, monkeypatch):
    """code_dir=None must not raise even though _save_code_artifact runs."""
    monkeypatch.setattr(
        codegen_path, "generate_slide_code",
        lambda *, layout_hint, pattern_guideline, slide_data, model: _VALID_SCRIPT,
    )
    out = tmp_path / "slide.pptx"
    codegen_path.render_via_codegen(
        _deck_meta(), _make_flat_slide(), out, code_dir=None,
    )
    assert out.exists()
