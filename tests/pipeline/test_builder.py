from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from src.pipeline.builder import build_presentation


HAPPY_CODE = '''
from pptx import Presentation
from pptx.util import Inches
import json, sys

def add_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text_frame.text = data["head_message"]

if __name__ == "__main__":
    data = json.loads(sys.stdin.read())
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_slide(prs, data)
    prs.save(sys.argv[1])
'''


@patch("src.pipeline.builder.revise_plan_until_pass", side_effect=lambda p, **kw: p)
@patch("src.pipeline.builder.make_plan")
@patch("src.pipeline.builder.generate_slide_code")
def test_build_falls_back_when_layout_hint_unknown(mock_codegen, mock_plan, mock_revise, tmp_path):
    mock_plan.return_value = {
        "deck_meta": {"title": "T", "theme": {}, "fonts": {}, "slide_size": {"width_in": 13.333, "height_in": 7.5}},
        "slides": [
            {"slide_no": 1, "purpose": "p", "head_message": "Hi", "layout_hint": "totally-fake-pattern", "content": {}},
        ],
    }
    mock_codegen.return_value = HAPPY_CODE
    out = tmp_path / "deck.pptx"
    build_presentation({"meta": {"title": "T"}, "sections": []},
                      output_path=out, workdir=tmp_path / "work")
    assert out.exists()


@patch("src.pipeline.builder.revise_plan_until_pass", side_effect=lambda p, **kw: p)
@patch("src.pipeline.builder.make_plan")
@patch("src.pipeline.builder.generate_slide_code")
def test_build_presentation_end_to_end(mock_codegen, mock_plan, mock_revise, tmp_path):
    mock_plan.return_value = {
        "deck_meta": {"title": "T", "theme": {"primary": "#065A82"}, "fonts": {}, "slide_size": {"width_in": 13.333, "height_in": 7.5}},
        "slides": [
            {"slide_no": 1, "purpose": "open", "head_message": "First", "layout_hint": "Cover", "content": {}},
            {"slide_no": 2, "purpose": "main", "head_message": "Second", "layout_hint": "Bullet List", "content": {}},
        ],
    }
    mock_codegen.return_value = HAPPY_CODE

    content = {"meta": {"title": "T"}, "sections": []}
    out = tmp_path / "deck.pptx"
    workdir = tmp_path / "work"

    build_presentation(content, output_path=out, workdir=workdir)

    assert out.exists()
    prs = Presentation(out)
    assert len(prs.slides) == 2
    # plan.json should be persisted in workdir for debugging
    assert (workdir / "plan.json").exists()
