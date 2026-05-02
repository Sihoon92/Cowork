from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from src.pipeline.text_critic import (
    extract_slide_text,
    critique_slide_text,
)


def _make_pptx(path, *texts):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = 0.5
    for t in texts:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(10), Inches(1))
        box.text_frame.text = t
        y += 1.0
    prs.save(path)
    return path


def test_extract_slide_text_returns_all_text_frames(tmp_path):
    p = _make_pptx(tmp_path / "x.pptx", "First", "Second", "Third")
    texts = extract_slide_text(p)
    assert "First" in texts
    assert "Second" in texts
    assert "Third" in texts


def test_extract_slide_text_skips_empty(tmp_path):
    p = _make_pptx(tmp_path / "x.pptx", "Real text", "")
    texts = extract_slide_text(p)
    assert "Real text" in texts
    assert "" not in texts


@patch("src.pipeline.text_critic.chat")
def test_critique_slide_text_pass(mock_chat, tmp_path):
    mock_chat.return_value = '''```json
{"issues": [], "verdict": "PASS"}
```'''
    p = _make_pptx(tmp_path / "x.pptx", "Hello world")
    slide_plan = {"slide_no": 1, "head_message": "Hello world", "layout_hint": "Cover", "content": {}}
    result = critique_slide_text(slide_plan, p)
    assert result["verdict"] == "PASS"
    assert result["issues"] == []


@patch("src.pipeline.text_critic.chat")
def test_critique_slide_text_fix(mock_chat, tmp_path):
    mock_chat.return_value = '''```json
{"issues": [{"severity":"high","msg":"Lorem ipsum placeholder found"}], "verdict": "FIX"}
```'''
    p = _make_pptx(tmp_path / "x.pptx", "Lorem ipsum dolor")
    slide_plan = {"slide_no": 2, "head_message": "Real content", "layout_hint": "Bullet List", "content": {}}
    result = critique_slide_text(slide_plan, p)
    assert result["verdict"] == "FIX"
    assert len(result["issues"]) == 1
    assert "Lorem" in result["issues"][0]["msg"]
