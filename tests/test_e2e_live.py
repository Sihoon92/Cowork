"""End-to-end test that hits real Ollama. Run with: pytest -m live"""
from pathlib import Path
import json

import pytest
from pptx import Presentation

from src.llm.ollama_client import is_available
from src.pipeline.builder import build_presentation


pytestmark = pytest.mark.live


@pytest.fixture
def sample_content():
    root = Path(__file__).resolve().parents[1]
    return json.loads((root / "data" / "sample_content.json").read_text(encoding="utf-8"))


def test_e2e_pipeline_produces_pptx(sample_content, tmp_path):
    if not is_available():
        pytest.skip("Ollama server not running")
    out = tmp_path / "demo.pptx"
    workdir = tmp_path / "work"
    build_presentation(sample_content, output_path=out, workdir=workdir)
    assert out.exists()
    prs = Presentation(out)
    assert len(prs.slides) >= 2
    plan = json.loads((workdir / "plan.json").read_text(encoding="utf-8"))
    assert "slides" in plan
    assert len(plan["slides"]) == len(prs.slides)
