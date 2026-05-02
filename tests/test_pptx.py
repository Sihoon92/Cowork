"""Test 2: JSON → PPTX 생성 테스트."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.pptx.generator import build_from_json

SAMPLE_JSON = Path(__file__).parent.parent / "data" / "sample_presentation.json"
OUTPUT_PATH = Path(__file__).parent.parent / "output" / "test_output.pptx"


def test_pptx_created():
    output = build_from_json(SAMPLE_JSON, OUTPUT_PATH)
    assert output.exists(), f"파일이 생성되지 않았습니다: {output}"
    size_kb = output.stat().st_size / 1024
    print(f"✅ PPTX 생성 성공: {output}")
    print(f"   파일 크기: {size_kb:.1f} KB")


def test_slide_count():
    from pptx import Presentation
    prs = Presentation(OUTPUT_PATH)
    assert len(prs.slides) == 3, f"슬라이드 수가 3이어야 하는데 {len(prs.slides)}입니다."
    print(f"✅ 슬라이드 수 확인: {len(prs.slides)}개")


def test_slide_dimensions():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(OUTPUT_PATH)
    assert abs(prs.slide_width - Inches(13.33)) < 1000, "슬라이드 너비 불일치"
    assert abs(prs.slide_height - Inches(7.5)) < 1000, "슬라이드 높이 불일치"
    print("✅ 슬라이드 크기 확인: 13.33 × 7.5 인치 (16:9)")


if __name__ == "__main__":
    print("=" * 50)
    print("Test 2: PPTX 생성 테스트")
    print("=" * 50)
    test_pptx_created()
    test_slide_count()
    test_slide_dimensions()
    print("\n모든 PPTX 테스트 통과!")
    print(f"\n생성된 파일: {OUTPUT_PATH.resolve()}")
