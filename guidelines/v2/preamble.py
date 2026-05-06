"""Prepended to every LLM-generated slide module. Provides imports,
design tokens, and a defense layer (`_patch_shapes`) against the
float-EMU bug python-pptx silently passes through.

The footer (added by code_assembler.py) calls _patch_shapes(slide)
before invoking build_slide(prs), so the LLM cannot disable the defense.
"""
from __future__ import annotations

import math
from pathlib import Path

# python-pptx
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

# matplotlib (Agg backend; Korean-capable font fallback)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.family"] = ["Pretendard", "Malgun Gothic", "sans-serif"]
plt.rcParams["axes.unicode_minus"] = False

# Color tokens
PRIMARY = RGBColor(0x1F, 0x36, 0x5C)
ACCENT  = RGBColor(0xE8, 0x6A, 0x33)
INK     = RGBColor(0x1A, 0x1A, 0x1A)
MUTED   = RGBColor(0x77, 0x77, 0x77)
BG      = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF4, 0xF6, 0xFA)

# Typography tokens
FONT_HEAD = "Pretendard"
FONT_BODY = "Pretendard"
SIZE_HEAD = Pt(36)
SIZE_SUB  = Pt(20)
SIZE_BODY = Pt(14)
SIZE_CAP  = Pt(11)

# Canvas (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _coerce_int_args(fn):
    def wrapper(*args, **kwargs):
        new_args = tuple(int(a) if isinstance(a, float) else a for a in args)
        new_kwargs = {
            k: (int(v) if isinstance(v, float) else v) for k, v in kwargs.items()
        }
        return fn(*new_args, **new_kwargs)
    wrapper.__wrapped__ = fn
    return wrapper


def _patch_shapes(slide):
    """Replace add_* methods on slide.shapes to coerce floats to ints.

    Safe to call multiple times — checks for prior wrapping.
    """
    for name in ("add_textbox", "add_shape", "add_connector",
                 "add_picture", "add_chart", "add_table"):
        orig = getattr(slide.shapes, name, None)
        if orig is None or hasattr(orig, "__wrapped__"):
            continue
        setattr(slide.shapes, name, _coerce_int_args(orig))
    return slide
